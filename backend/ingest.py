"""Review-file ingestion and validation for CSV and XLSX uploads.

Scope: parse, validate, pick the review-text column, normalize and redact.
No NLP happens here. Uploaded review text lives only in memory and is never
logged or written to disk by this module.

Conventions
- "Rows" are data rows below the header, ignoring completely blank rows.
- The row cap is enforced while reading, before detection, cleaning or
  redaction, so oversized files are rejected cheaply.
- Spreadsheet formulas are never evaluated. XLSX is read with cached values
  only; CSV cells are plain text. Downstream exports must still neutralize
  leading "=", "+", "-", "@" characters before writing CSV/XLSX.
"""

from __future__ import annotations

import csv
import io
import logging
import os
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime
from xml.etree.ElementTree import ParseError

import pandas as pd
from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from redaction import redact_pii
from schemas import IngestMetadata

logger = logging.getLogger("insightforge.ingest")

# --- Limits that are not (yet) configurable ---------------------------------
MAX_COLUMNS = 100
MAX_SCAN_ROWS = 100_000  # physical XLSX rows scanned, blank or not
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024  # zip-bomb guard
MAX_XLSX_ENTRIES = 1_000
MAX_NAME_CHARS = 80
MAX_PICKER_COLUMNS = 50
_READ_CHUNK = 64 * 1024

# --- Column aliases (compared to normalized header names) -------------------
TEXT_ALIASES = frozenset(
    {
        "review_text", "reviewtext", "review_body", "review_content",
        "review_comment", "review_comments", "review", "reviews",
        "customer_review", "customer_reviews", "user_review", "text",
        "comment", "comments", "feedback", "customer_feedback", "body",
        "content",
    }
)  # fmt: skip
RATING_ALIASES = frozenset(
    {
        "rating", "ratings", "star", "stars", "star_rating", "review_rating",
        "overall_rating", "overall",
    }
)  # fmt: skip
DATE_ALIASES = frozenset(
    {
        "date", "review_date", "reviewed_at", "review_time", "created_at",
        "created", "timestamp", "posted_on", "posted_at", "date_posted", "time",
    }
)  # fmt: skip
VERSION_ALIASES = frozenset(
    {
        "version", "app_version", "product_version", "review_version",
        "appversion", "software_version", "build",
    }
)  # fmt: skip

_MISSING_TOKENS = frozenset({"", "nan", "none", "null", "n/a", "#n/a"})
_NUMERIC_RE = re.compile(
    r"[+-]?(?:\d[\d,]*(?:\.\d+)?|\.\d+)%?|\d+(?:\.\d+)?/\d+"
)
_WHITESPACE_RE = re.compile(r"\s+")
# Control characters and zero-width space/BOM. U+200C/U+200D (ZWNJ/ZWJ) are
# deliberately kept: Indic scripts rely on them.
_DROP_CHARS = {
    ord(c): None for c in [*map(chr, range(0, 32)), *map(chr, range(0x7F, 0xA0)), "\u200b", "\ufeff"]
}  # fmt: skip

_SUPPORTED_EXTENSIONS = {".csv": "csv", ".xlsx": "xlsx"}
# openpyxl signals malformed packages with these; input is an in-memory buffer,
# so OSError here means "bad workbook", not a real I/O failure.
_XLSX_ERRORS = (
    zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, ParseError, EOFError, OSError,
)  # fmt: skip


class IngestError(Exception):
    """Expected, user-facing ingestion failure. Messages never contain review text."""

    def __init__(
        self,
        code: str,
        message: str,
        status_code: int = 400,
        columns: list[str] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code
        self.columns = columns


@dataclass(frozen=True)
class Column:
    index: int
    original: str  # display name, de-duplicated
    base: str  # normalized name used for alias matching
    key: str  # unique normalized name for internal use


@dataclass
class RawTable:
    headers: list[str]
    rows: list[list[str]]
    row_numbers: list[int]
    encoding: str | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass
class IngestResult:
    reviews: pd.DataFrame  # source_row, text, text_redacted, + optional columns
    metadata: IngestMetadata


# --- Upload reading and file-type checks ------------------------------------
def _format_bytes(n: int) -> str:
    return f"{n / (1024 * 1024):.1f} MB" if n >= 1024 * 1024 else f"{max(n // 1024, 1)} KB"


def file_too_large(max_bytes: int) -> IngestError:
    return IngestError(
        "file_too_large", f"The file is larger than the {_format_bytes(max_bytes)} limit.", 413
    )


async def read_upload(file: UploadFile, max_bytes: int) -> bytes:
    """Read an upload in chunks, aborting as soon as it exceeds ``max_bytes``."""
    buffer = bytearray()
    while chunk := await file.read(_READ_CHUNK):
        buffer.extend(chunk)
        if len(buffer) > max_bytes:
            raise file_too_large(max_bytes)
    return bytes(buffer)


def sanitize_filename(filename: str | None) -> str:
    name = (filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    name = "".join(ch for ch in name if ch.isprintable()).strip()
    return name[:100]


def detect_file_kind(filename: str, data: bytes) -> str:
    """Return "csv" or "xlsx" after checking extension AND content signature."""
    kind = _SUPPORTED_EXTENSIONS.get(os.path.splitext(filename)[1].lower())
    if kind is None:
        raise IngestError(
            "unsupported_file_type", "Unsupported file type. Upload a .csv or .xlsx file.", 415
        )
    if not data:
        raise IngestError("empty_file", "The uploaded file is empty.")
    is_zip = data.startswith(b"PK\x03\x04")
    if kind == "xlsx" and not is_zip:
        hint = " Legacy .xls and password-protected workbooks are not supported."
        raise IngestError(
            "invalid_file", "The file is not a valid .xlsx workbook." + hint
        )
    if kind == "csv" and (is_zip or data.startswith(b"\xd0\xcf\x11\xe0")):
        raise IngestError("invalid_file", "The file content does not look like CSV text.")
    return kind


# --- Parsing ----------------------------------------------------------------
def decode_csv_bytes(data: bytes) -> tuple[str, str]:
    """Decode bytes, trying UTF-8 (with BOM), UTF-16 (with BOM), then fallbacks."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            raise IngestError("invalid_file", "The file could not be decoded as text.") from None
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig", errors="replace"), "utf-8-sig"
    for encoding in ("utf-8", "cp1252"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), "latin-1"  # never fails


def _looks_binary(text: str) -> bool:
    sample = text[:8192]
    if not sample:
        return False
    if "\x00" in sample:
        return True
    controls = sum(1 for ch in sample if ch not in "\t\r\n" and unicodedata.category(ch) == "Cc")
    return controls / len(sample) > 0.05


def _sniff_delimiter(text: str) -> str:
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _csv_records(text: str) -> Iterator[tuple[int, list[str]]]:
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_sniff_delimiter(text), strict=True)
    for record in reader:
        yield reader.line_num, record


def _is_blank(record: Iterable[str]) -> bool:
    return all(not cell.strip() for cell in record)


def _collect_rows(records: Iterable[tuple[int, list[str]]], max_rows: int) -> RawTable:
    """Shared header/row collection with the row cap enforced during reading."""
    headers: list[str] | None = None
    rows: list[list[str]] = []
    numbers: list[int] = []
    for number, record in records:
        if _is_blank(record):
            continue
        if headers is None:
            headers = list(record)
            while headers and not headers[-1].strip():
                headers.pop()
            if len(headers) > MAX_COLUMNS:
                raise IngestError("invalid_file", f"The file has more than {MAX_COLUMNS} columns.")
            continue
        if len(rows) >= max_rows:
            raise IngestError(
                "too_many_rows", f"The file has more than {max_rows:,} rows. Split it and retry.", 422
            )
        extra = record[len(headers):]
        if any(cell.strip() for cell in extra):
            raise IngestError(
                "invalid_file", f"Row {number} has more fields than the header row."
            )
        rows.append(list(record[: len(headers)]) + [""] * (len(headers) - len(record)))
        numbers.append(number)
    if headers is None:
        raise IngestError("empty_file", "The uploaded file has no content.")
    return RawTable(headers, rows, numbers)


def parse_csv(data: bytes, max_rows: int) -> RawTable:
    text, encoding = decode_csv_bytes(data)
    if _looks_binary(text):
        raise IngestError("invalid_file", "The file does not look like a text CSV file.")
    if not text.strip():
        raise IngestError("empty_file", "The uploaded file has no content.")
    try:
        table = _collect_rows(_csv_records(text), max_rows)
    except csv.Error:
        raise IngestError(
            "invalid_file", "The CSV is malformed. Check quotes and delimiters."
        ) from None
    table.encoding = encoding
    if encoding not in ("utf-8", "utf-8-sig", "utf-16"):
        table.warnings.append(
            f"The file was not valid UTF-8 and was read as {encoding}. "
            "Some characters may be incorrect; re-export as UTF-8 if so."
        )
    return table


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _check_xlsx_container(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            names = {info.filename for info in infos}
    except zipfile.BadZipFile:
        raise IngestError("invalid_file", "The .xlsx file is corrupt or unreadable.") from None
    if "xl/workbook.xml" not in names:
        raise IngestError("invalid_file", "The file is not a valid .xlsx workbook.")
    if len(infos) > MAX_XLSX_ENTRIES or sum(i.file_size for i in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise IngestError("invalid_file", "The .xlsx file is too large once uncompressed.")


def _xlsx_records(worksheet) -> Iterator[tuple[int, list[str]]]:
    rows = worksheet.iter_rows(values_only=True, max_col=MAX_COLUMNS + 1)
    for number, record in enumerate(rows, start=1):
        if number > MAX_SCAN_ROWS:
            raise IngestError("invalid_file", "The worksheet has too many rows to scan.")
        yield number, [_cell_to_str(v) for v in record]


def parse_xlsx(data: bytes, max_rows: int) -> RawTable:
    """Read the first worksheet using cached cell values (formulas never run)."""
    _check_xlsx_container(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except _XLSX_ERRORS:
        raise IngestError("invalid_file", "The .xlsx file is corrupt or unreadable.") from None
    try:
        if not workbook.worksheets:
            raise IngestError("invalid_file", "The workbook has no worksheets.")
        table = _collect_rows(_xlsx_records(workbook.worksheets[0]), max_rows)
        if len(workbook.worksheets) > 1:
            table.warnings.append("Only the first worksheet was read.")
    except _XLSX_ERRORS:
        raise IngestError("invalid_file", "The .xlsx file is corrupt or unreadable.") from None
    finally:
        workbook.close()
    return table


# --- Normalization ----------------------------------------------------------
def clean_text(value: str) -> str:
    """Collapse whitespace, drop control characters, map missing tokens to ''."""
    text = unicodedata.normalize("NFC", value)
    text = _WHITESPACE_RE.sub(" ", text).translate(_DROP_CHARS).strip()
    return "" if text.lower() in _MISSING_TOKENS else text


def normalize_column_name(name: str) -> str:
    text = unicodedata.normalize("NFKC", name).casefold()
    return re.sub(r"\W+", "_", text).strip("_")


def _safe_name(name: str) -> str:
    return "".join(ch for ch in name if ch.isprintable())[:MAX_NAME_CHARS]


def _safe_names(names: Iterable[str]) -> list[str]:
    return [_safe_name(n) for n in list(names)[:MAX_PICKER_COLUMNS]]


def build_columns(headers: list[str]) -> tuple[list[Column], list[str]]:
    """Give every column a unique display name and a safe normalized key."""
    columns: list[Column] = []
    warnings: list[str] = []
    seen_originals: dict[str, int] = {}
    seen_keys: dict[str, int] = {}
    for index, raw in enumerate(headers):
        original = _WHITESPACE_RE.sub(" ", raw).strip() or f"Unnamed column {index + 1}"
        base = normalize_column_name(original) or f"column_{index + 1}"
        count = seen_originals.get(original, 0) + 1
        seen_originals[original] = count
        if count > 1:
            warnings.append("Duplicate column names were renamed.")
            original = f"{original} ({count})"
        key_count = seen_keys.get(base, 0) + 1
        seen_keys[base] = key_count
        key = base if key_count == 1 else f"{base}_{key_count}"
        columns.append(Column(index, original, base, key))
    return columns, sorted(set(warnings))


# --- Column detection -------------------------------------------------------
def _values(table: RawTable, column: Column) -> list[str]:
    return [clean_text(row[column.index]) for row in table.rows]


def _is_text_like(values: list[str]) -> bool:
    filled = [v for v in values if v]
    if not filled:
        return False
    numeric = sum(1 for v in filled if _NUMERIC_RE.fullmatch(v))
    return numeric / len(filled) < 0.8


def _picker_candidates(columns: list[Column], table: RawTable) -> list[str]:
    optional = RATING_ALIASES | DATE_ALIASES | VERSION_ALIASES
    return _safe_names(
        c.original for c in columns if c.base not in optional and _is_text_like(_values(table, c))
    )


def select_text_column(columns: list[Column], table: RawTable, requested: str | None) -> Column:
    if requested and requested.strip():
        return _resolve_requested(columns, table, requested.strip())
    usable = [
        c for c in columns if c.base in TEXT_ALIASES and _is_text_like(_values(table, c))
    ]
    if len(usable) == 1:
        return usable[0]
    if len(usable) > 1:
        raise IngestError(
            "ambiguous_text_column",
            "Several columns could contain the review text. Choose one.",
            422,
            _safe_names(c.original for c in usable),
        )
    if len(columns) == 1 and _is_text_like(_values(table, columns[0])):
        return columns[0]
    aliased = [c for c in columns if c.base in TEXT_ALIASES]
    if len(aliased) == 1 and not any(_values(table, aliased[0])):
        return aliased[0]  # named like a review column but blank: report empty_dataset
    raise IngestError(
        "no_text_column",
        "No review-text column was found. Choose the column that contains the reviews.",
        422,
        _picker_candidates(columns, table),
    )


def _resolve_requested(columns: list[Column], table: RawTable, wanted: str) -> Column:
    matches = [c for c in columns if c.original == wanted]
    if not matches:
        normalized = normalize_column_name(wanted)
        matches = [c for c in columns if normalized and c.base == normalized]
    if len(matches) > 1:
        raise IngestError(
            "ambiguous_text_column",
            "The selected column name matches several columns.",
            422,
            _safe_names(c.original for c in matches),
        )
    if not matches:
        raise IngestError(
            "invalid_text_column",
            "The selected column was not found in the file.",
            422,
            _picker_candidates(columns, table),
        )
    if not _is_text_like(_values(table, matches[0])):
        raise IngestError(
            "invalid_text_column",
            "The selected column does not contain usable review text.",
            422,
            _picker_candidates(columns, table),
        )
    return matches[0]


def _find_optional(columns: list[Column], aliases: frozenset[str], taken: set[int]) -> Column | None:
    for column in columns:
        if column.index not in taken and column.base in aliases:
            return column
    return None


def coerce_and_validate_rating(raw: str | None) -> tuple[float | None, str | None, bool, bool]:
    """Parse, coerce to 1-5, and validate numeric rating per PRD v4 Section 11.

    Returns (rating, rating_source, is_out_of_range, is_invalid).
    """
    if raw is None or not str(raw).strip():
        return None, None, False, False
    s = str(raw).strip()
    if s.lower() in _MISSING_TOKENS:
        return None, None, False, False

    num: float | None = None
    # Check "X/5"
    m_5 = re.match(r"^(\d+(?:\.\d+)?)\s*/\s*5(?:\.0)?$", s)
    if m_5:
        num = float(m_5.group(1))
    else:
        # Check "X/10"
        m_10 = re.match(r"^(\d+(?:\.\d+)?)\s*/\s*10(?:\.0)?$", s)
        if m_10:
            num = float(m_10.group(1)) / 2.0
        else:
            # Check "X star(s)"
            m_star = re.match(r"^(\d+(?:\.\d+)?)\s*stars?$", s, re.IGNORECASE)
            if m_star:
                num = float(m_star.group(1))
            else:
                try:
                    num = float(s)
                except ValueError:
                    # Non-numeric string: drop invalid rating per PRD
                    return None, None, False, True

    if num < 1.0 or num > 5.0:
        # Coerce to 1 to 5 per PRD Section 11
        coerced = max(1.0, min(5.0, round(num, 1)))
        return coerced, "given", True, False

    return round(num, 1), "given", False, False


# --- Frame building and orchestration ---------------------------------------
def build_reviews_frame(
    table: RawTable, text_col: Column, optional: dict[str, Column | None]
) -> tuple[pd.DataFrame, int, list[str]]:
    """Keep only needed columns; exclude rows whose review text is empty.
    Flag exact duplicates and validate/coerce ratings per PRD.
    """
    names = [name for name, col in optional.items() if col is not None]
    records: list[dict[str, object]] = []
    excluded = 0
    seen_texts: set[str] = set()
    dup_count = 0
    invalid_rating_count = 0
    out_of_range_rating_count = 0
    has_rating = "rating" in optional and optional["rating"] is not None

    for number, row in zip(table.row_numbers, table.rows):
        text = clean_text(row[text_col.index])
        if not text:
            excluded += 1
            continue

        is_dup = text in seen_texts
        if is_dup:
            dup_count += 1
        else:
            seen_texts.add(text)

        record: dict[str, object] = {
            "source_row": number,
            "text": text,
            "text_redacted": redact_pii(text),
            "is_duplicate": is_dup,
        }

        if has_rating:
            raw_rating = clean_text(row[optional["rating"].index]) or None  # type: ignore[union-attr]
            r_val, r_src, out_range, inv = coerce_and_validate_rating(raw_rating)
            if out_range:
                out_of_range_rating_count += 1
            if inv:
                invalid_rating_count += 1
            record["rating"] = r_val
            record["rating_source"] = r_src
        else:
            record["rating_source"] = None

        for name in names:
            if name != "rating":
                record[name] = clean_text(row[optional[name].index]) or None  # type: ignore[union-attr]
        records.append(record)

    extra_warnings: list[str] = []
    if dup_count > 0:
        extra_warnings.append(f"Flagged {dup_count} exact duplicate review(s).")
    if out_of_range_rating_count > 0:
        extra_warnings.append(f"Coerced {out_of_range_rating_count} rating(s) outside 1-5 range.")
    if invalid_rating_count > 0:
        extra_warnings.append(f"Dropped {invalid_rating_count} invalid non-numeric rating(s).")

    col_order = ["source_row", "text", "text_redacted", "is_duplicate", "rating_source", *names]
    frame = pd.DataFrame.from_records(records, columns=col_order)
    return frame, excluded, extra_warnings


def ingest_file(
    filename: str, data: bytes, *, max_rows: int, text_column: str | None = None
) -> IngestResult:
    """Validate and parse an uploaded review file. Raises IngestError on bad input."""
    kind = detect_file_kind(filename, data)
    table = parse_csv(data, max_rows) if kind == "csv" else parse_xlsx(data, max_rows)
    if not table.rows:
        raise IngestError("empty_dataset", "The file has a header row but no data rows.", 422)

    columns, column_warnings = build_columns(table.headers)
    text_col = select_text_column(columns, table, text_column)
    taken = {text_col.index}
    optional: dict[str, Column | None] = {}
    for name, aliases in (("rating", RATING_ALIASES), ("date", DATE_ALIASES), ("version", VERSION_ALIASES)):
        optional[name] = _find_optional(columns, aliases, taken)
        if optional[name] is not None:
            taken.add(optional[name].index)  # type: ignore[union-attr]

    frame, excluded, ingestion_warnings = build_reviews_frame(table, text_col, optional)
    if frame.empty:
        raise IngestError("empty_dataset", "Every row has an empty review text.", 422)

    def original(name: str) -> str | None:
        col = optional[name]
        return _safe_name(col.original) if col else None

    metadata = IngestMetadata(
        source_name=sanitize_filename(filename),
        file_type=kind,  # type: ignore[arg-type]
        encoding=table.encoding,
        original_columns=_safe_names(c.original for c in columns),
        text_column=_safe_name(text_col.original),
        rating_column=original("rating"),
        date_column=original("date"),
        version_column=original("version"),
        rows_read=len(table.rows),
        excluded_empty=excluded,
        usable_rows=len(frame),
        warnings=[*table.warnings, *column_warnings, *ingestion_warnings],
    )
    logger.info(
        "ingest ok type=%s rows=%d excluded=%d", kind, metadata.rows_read, excluded
    )
    return IngestResult(frame, metadata)
